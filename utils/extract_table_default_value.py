import docx
import openpyxl
from docx import Document
from openpyxl.styles import Font

from utils.PathManager import load_path_manager as lpm

file_refer = lpm.refer("Program Specification.docx")
read_file = lpm.input("ping.xlsx")
fill_file = lpm.input("fill-ping.xlsx")


def extract_table_value():
    doc = Document(file_refer)

    api_sections = []
    api_section_data = {}
    current_section = None
    table_count = len(doc.tables)  # Note
    # print(table_count)

    for i, element in enumerate(doc.element.body):
        if isinstance(element, docx.oxml.text.paragraph.CT_P):
            paragraph = docx.text.paragraph.Paragraph(element, doc)
            text = paragraph.text.strip()

            if text.startswith("AA-") or text.startswith("BB-"):
                if current_section:
                    api_sections.append(api_section_data)
                    api_section_data = {}

                current_section = text
                api_section_data["API Section"] = current_section
                print("----------------------------------------------------------")
                print(current_section)
            elif text.startswith("API ID"):
                strip = text.split(":")[-1].strip()
                if strip not in current_section:
                    print("not matching API ID")
                    api_section_data["API ID"] = current_section[:12]  # handle prog spec API ID typo
                else:
                    api_section_data["API ID"] = text.split(":")[-1].strip()

                # print(text.split(":")[-1].strip())

            elif text.startswith("Path Parameter(s)"):
                print(text)
            elif text.startswith("Request Body"):
                print(text)
            elif text.startswith("Data Access Layer"):
                print(text)
            elif text.startswith("Operation:"):
                print(text)
                # if "Data Access Layer Operation" not in api_section_data:
                #     api_section_data["Data Access Layer Operation"] = text.replace("Operation:", "").strip()
                # else:
                #     api_section_data["Data Access Layer Operation"] += "," + text.replace("Operation:", "").strip()

        elif isinstance(element, docx.oxml.table.CT_Tbl):
            table = docx.table.Table(element, doc)
            # print("ping table")
            # print(len(table.rows))

            first_row = table.rows[0]
            if len(first_row.cells) > 1:
                first_cell_text = first_row.cells[0].text.strip()
                second_cell_text = first_row.cells[1].text.strip()

                header_row = table.rows[0]
                headers = [cell.text.strip() for cell in header_row.cells]
                # column_data = {header: [] for header in headers}
                column_data_row = []
                if (first_cell_text, second_cell_text) in [("Parameters Name", "Value"),
                                                           ("Parameters Name", "Possible Values"),
                                                           ("Source Table", "Source Field Name"),
                                                           ("Source", "Source Field Name")]:
                    key = f"{first_cell_text}-{second_cell_text}"

                    # for row in table.rows[1:]:
                    #     for header, cell in zip(headers, row.cells):
                    #         column_data[header].append(cell.text.strip())

                    # column_data_row.append(headers)

                    for row in table.rows[1:]:
                        row_data = []
                        for cell in row.cells:
                            row_data.append(cell.text.strip())
                        column_data_row.append(row_data)

                    # print("Headers:", headers)
                    # for header, data in column_data.items():
                    #     print(f"{header}: {data}")

                    if key == "Parameters Name-Value":
                        api_section_data["IP-PP-Default-Value"] = column_data_row
                    elif key == "Parameters Name-Possible Values":
                        api_section_data["IP-RB-Default-Value"] = column_data_row
                    elif key == "Source Table-Source Field Name":
                        api_section_data["BRL-DAL-R-Default-Value"] = column_data_row
                    elif key == "Source-Source Field Name":
                        api_section_data["BRL-DAL-W-Default-Value"] = column_data_row

        else:
            continue

    if current_section:
        api_sections.append(api_section_data)

    # print(api_sections)
    return api_sections


def clear_table_value(tb_data):
    reorganized_data = {
        'IP-PP-Default-Value': {},
        'IP-RB-Default-Value': {},
        'BRL-DAL-R-Default-Value': {},
        'BRL-DAL-W-Default-Value': {}
    }

    for item in tb_data:
        for key in reorganized_data.keys():
            if key in item:
                reorganized_data[key].update({item['API ID']: item[key]})

    # print(reorganized_data)
    return reorganized_data


def fill_data_to_master_data_excel(inc_data):
    print("fill data")
    workbook = openpyxl.load_workbook(read_file)

    count_ip_pp = 0
    count_ip_rb = 0
    count_dal_r = 0
    count_dal_w = 0
    for key, values in inc_data.items():
        sheet = workbook[key]

        print(f"fill {key} sheet data")
        # distinct
        distinct_check = []
        for item_id, item_data in values.items():

            for i_i_data in item_data:

                if key == "IP-PP-Default-Value":

                    if i_i_data[0] not in distinct_check:
                        sheet.append([i_i_data[0], i_i_data[2]])
                        count_ip_pp += 1
                        distinct_check.append(i_i_data[0])

                if key == "IP-RB-Default-Value":

                    if i_i_data[0] not in distinct_check:
                        sheet.append([i_i_data[0], i_i_data[1], i_i_data[2]])
                        count_ip_rb += 1
                        distinct_check.append(i_i_data[0])

                if key == "BRL-DAL-R-Default-Value":
                    exists = False

                    if distinct_check:
                        for item in distinct_check:
                            if [i_i_data[0], i_i_data[3]] == item:
                                exists = True
                                break

                    if exists is False:
                        sheet.append([i_i_data[0], i_i_data[3], i_i_data[4]])
                        count_dal_r += 1
                        distinct_check.append([i_i_data[0], i_i_data[3]])

                    # else:
                    #     sheet.append([i_i_data[0], i_i_data[3], i_i_data[4]])
                    #     distinct_check.append([i_i_data[0], i_i_data[3]])

                if key == "BRL-DAL-W-Default-Value":
                    exists = False

                    if distinct_check:
                        for item in distinct_check:
                            if [i_i_data[1], i_i_data[2]] == item:
                                exists = True
                                break

                                # sheet.append([i_i_data[1], i_i_data[2], i_i_data[4]])
                                # count_dal_w += 1
                                # distinct_check.append([i_i_data[1], i_i_data[2]])
                                # break
                    # else:
                    #     sheet.append([i_i_data[1], i_i_data[2], i_i_data[4]])
                    #     distinct_check.append([i_i_data[1], i_i_data[2]])

                    if exists is False:
                        sheet.append([i_i_data[1], i_i_data[2], i_i_data[4]])
                        count_dal_w += 1
                        distinct_check.append([i_i_data[1], i_i_data[2]])

    workbook.save(fill_file)
    print("finish")
    print(f"finish count_ip_pp {count_ip_pp}")
    print(f"finish count_ip_rb {count_ip_rb}")
    print(f"finish count_dal_r {count_dal_r}")
    print(f"finish count_dal_w {count_dal_w}")


if __name__ == '__main__':
    table_value = extract_table_value()
    compose_data = clear_table_value(table_value)
    fill_data_to_master_data_excel(compose_data)
